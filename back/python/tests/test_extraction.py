from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document

from app.extraction import UnsupportedFile, extract_resume


def test_txt_offsets_are_stable():
    result = extract_resume("TXT", "alpha\nbeta".encode())
    assert result.text == "alpha\nbeta"
    assert (result.evidence[1]["sourceStart"], result.evidence[1]["sourceEnd"]) == (6, 10)


def test_docx_paragraph_offsets_are_stable():
    document = Document()
    document.add_paragraph("alpha")
    document.add_paragraph("beta")
    stream = BytesIO()
    document.save(stream)
    result = extract_resume("DOCX", stream.getvalue())
    assert result.text == "alpha\nbeta"
    assert result.evidence[1]["sourceStart"] == 6


def test_docx_table_paragraphs_are_extracted_in_document_order():
    document = Document()
    document.add_paragraph("outside first")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "inside table"
    document.add_paragraph("outside final")
    stream = BytesIO()
    document.save(stream)

    result = extract_resume("DOCX", stream.getvalue())

    assert result.text == "outside first\ninside table\noutside final"
    assert [item["sourceLocation"] for item in result.evidence] == [
        "paragraph:0",
        "paragraph:1",
        "paragraph:2",
    ]
    assert [(item["sourceStart"], item["sourceEnd"]) for item in result.evidence] == [
        (0, 13),
        (14, 26),
        (27, 40),
    ]


def test_docx_soft_break_stays_in_its_physical_paragraph_evidence_range():
    document = Document()
    first = document.add_paragraph()
    first.add_run("alpha").add_break()
    first.add_run("beta")
    document.add_paragraph("gamma")
    stream = BytesIO()
    document.save(stream)

    result = extract_resume("DOCX", stream.getvalue())

    assert result.text == "alpha\nbeta\ngamma"
    assert [item["sourceLocation"] for item in result.evidence] == [
        "paragraph:0",
        "paragraph:1",
    ]
    assert [(item["sourceStart"], item["sourceEnd"]) for item in result.evidence] == [
        (0, 10),
        (11, 16),
    ]
    assert [item["excerpt"] for item in result.evidence] == ["alpha\nbeta", "gamma"]


def test_other_types_are_unsupported():
    with pytest.raises(UnsupportedFile):
        extract_resume("PDF", b"%PDF")


def test_malformed_docx_is_unsupported():
    with pytest.raises(UnsupportedFile):
        extract_resume("DOCX", b"not a docx package")


def test_incomplete_docx_package_is_unsupported():
    stream = BytesIO()
    with ZipFile(stream, "w") as package:
        package.writestr("word/document.xml", b"<broken>")
    with pytest.raises(UnsupportedFile):
        extract_resume("DOCX", stream.getvalue())
